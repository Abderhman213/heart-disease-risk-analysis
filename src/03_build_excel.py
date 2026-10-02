"""Exports the exact datasets behind the dashboard to Excel (data only, no charts)."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
clean=pd.read_csv("data/heart_clean.csv"); raw=pd.read_csv("data/heart_raw_combined.csv")
cd=pd.DataFrame({"patient_id":range(1,len(clean)+1),"site":clean.site,"age":clean.age,"age_group":clean.age_group,"sex":clean.sex_l,"chest_pain_type":clean.cp_l,
 "resting_bp":clean.trestbps,"cholesterol":clean.chol,"fasting_sugar_gt120":clean.fbs,"rest_ecg":clean.restecg,"max_heart_rate":clean.thalach,
 "exercise_angina":clean.exang_l,"st_depression":clean.oldpeak,"slope":clean.slope,"major_vessels_ca":clean.ca,"thal":clean.thal,
 "severity_0_4":clean.num,"heart_disease":clean.disease})
rd=raw.rename(columns={"trestbps":"resting_bp","chol":"cholesterol","fbs":"fasting_sugar_gt120","restecg":"rest_ecg","thalach":"max_heart_rate","exang":"exercise_angina_01","oldpeak":"st_depression","ca":"major_vessels_ca","num":"severity_0_4"}); rd.insert(0,"patient_id",range(1,len(rd)+1))
wb=Workbook(); wb.remove(wb.active)
def sheet(name,df,tname):
    ws=wb.create_sheet(name); ws.append(list(df.columns))
    for r in df.itertuples(index=False): ws.append([None if pd.isna(v) else (v.item() if hasattr(v,"item") else v) for v in r])
    ref=f"A1:{get_column_letter(df.shape[1])}{len(df)+1}"
    t=Table(displayName=tname,ref=ref); t.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True); ws.add_table(t)
    for i,c in enumerate(df.columns,1): ws.column_dimensions[get_column_letter(i)].width=max(12,min(24,len(c)+4))
    ws.freeze_panes="A2"
sheet("Clean Data (dashboard)",cd,"CleanData"); sheet("Raw Data (UCI combined)",rd,"RawData")
g=wb.create_sheet("Field Guide")
rows=[("Field","Meaning"),("site","Hospital: Cleveland, Hungary, Switzerland, VA Long Beach"),("age / age_group","Years; bands <40, 40-49, 50-59, 60-69, 70+"),("sex","Male / Female"),
("chest_pain_type","Typical angina, Atypical angina, Non-anginal pain, Asymptomatic"),("resting_bp","Resting blood pressure (mm Hg); impossible 0 set to blank"),("cholesterol","Serum cholesterol (mg/dl); impossible 0 set to blank (172 rows)"),
("fasting_sugar_gt120","1 if fasting blood sugar > 120 mg/dl"),("rest_ecg","0 normal, 1 ST-T abnormality, 2 LV hypertrophy"),("max_heart_rate","Maximum heart rate achieved in exercise test"),("exercise_angina","Chest pain induced by exercise"),
("st_depression","ST depression induced by exercise relative to rest"),("slope / major_vessels_ca / thal","Mostly missing, excluded from the model"),("severity_0_4","Angiographic vessel narrowing; 0 = none"),("heart_disease","1 if severity > 0 (dashboard target)"),
("",""),("Source","UCI Machine Learning Repository, Heart Disease (Detrano et al., 1989): archive.ics.uci.edu/dataset/45/heart+disease"),("Blank cells","Missing values")]
for r in rows: g.append(r)
for c in g[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="0F172A")
g.column_dimensions["A"].width=32; g.column_dimensions["B"].width=95
wb.save("outputs/Heart_Disease_Data.xlsx")
